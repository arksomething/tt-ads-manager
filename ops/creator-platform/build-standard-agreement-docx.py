#!/usr/bin/env python3
"""Build the public, non-binding GoTall creator-agreement sample DOCX."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCE = REPO_ROOT / "creator-platform" / "src" / "content" / "standard-creator-agreement-sample.json"
DEFAULT_OUTPUT = (
    REPO_ROOT
    / "creator-platform"
    / "public"
    / "documents"
    / "gotall-standard-creator-agreement-sample-v0.1.docx"
)

INK = "1F2328"
MUTED = "57606A"
LINE = "D0D7DE"
SOFT = "F6F8FA"
GREEN = "1F883D"
DARK = "24292F"
AMBER = "9A6700"
AMBER_BG = "FFF8C5"
WHITE = "FFFFFF"
FONT = "Liberation Sans"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge in edges.items():
        tag = f"w:{edge_name}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge.items():
            element.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name=FONT) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, text: str, url: str, color="0969DA", underline=True):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    run_properties.append(color_node)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        run_properties.append(underline_node)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    run_properties.append(fonts)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def set_paragraph_rules(paragraph, keep_next=False, keep_lines=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))
    p_pr.append(OxmlElement("w:widowControl"))


def add_text(paragraph, text: str, *, bold=False, color=INK, size=9.3, italic=False) -> None:
    run = paragraph.add_run(text)
    set_run_font(run)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Title", 28, INK, 0, 9),
        ("Heading 1", 15, INK, 15, 7),
        ("Heading 2", 11.5, INK, 10, 4),
        ("Heading 3", 10, MUTED, 8, 3),
    ):
        style = document.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    document.styles["List Bullet"].font.name = FONT
    document.styles["List Bullet"].font.size = Pt(9.1)
    document.styles["List Bullet"].paragraph_format.left_indent = Inches(0.23)
    document.styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.14)
    document.styles["List Bullet"].paragraph_format.space_after = Pt(3)


def add_header_footer(document: Document, data: dict) -> None:
    document.settings.odd_and_even_pages_header_footer = True

    def fill_header(header) -> None:
        table = header.add_table(rows=1, cols=2, width=Inches(7.06))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width = Inches(3.53)
        table.columns[1].width = Inches(3.53)
        left, right = table.rows[0].cells
        left_p = left.paragraphs[0]
        add_text(left_p, "GOTALL", bold=True, color=INK, size=8)
        add_text(left_p, "  CREATOR PROGRAM", bold=True, color=MUTED, size=7.3)
        right_p = right.paragraphs[0]
        right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_text(right_p, data["status"], bold=True, color=AMBER, size=7.3)
        for cell in (left, right):
            set_cell_margins(cell, top=0, start=0, bottom=50, end=0)
            set_cell_border(cell, bottom={"val": "single", "sz": 6, "color": LINE})

    def fill_footer(footer) -> None:
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, f'{data["version"]}  ·  Page ', color=MUTED, size=7.5)
        add_page_field(p)
        add_text(p, "  ·  Assigned signed version controls", color=MUTED, size=7.5)

    for section in document.sections:
        fill_header(section.header)
        fill_header(section.even_page_header)
        fill_footer(section.footer)
        fill_footer(section.even_page_footer)


def add_label(paragraph, text: str, *, color=MUTED) -> None:
    add_text(paragraph, text.upper(), bold=True, color=color, size=7.5)
    paragraph.paragraph_format.space_after = Pt(4)
    set_paragraph_rules(paragraph, keep_next=True)


def add_title_block(document: Document, data: dict) -> None:
    label = document.add_paragraph()
    add_label(label, "GoTall creator program")

    title = document.add_paragraph(style="Title")
    add_text(title, data["documentTitle"], bold=True, color=INK, size=28)
    set_paragraph_rules(title, keep_next=True)

    metadata = document.add_paragraph()
    add_text(metadata, f'{data["version"]}  ·  {data["documentDate"]}', color=MUTED, size=9)
    metadata.paragraph_format.space_after = Pt(12)

    status_table = document.add_table(rows=1, cols=1)
    status_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    status_cell = status_table.cell(0, 0)
    set_cell_shading(status_cell, AMBER_BG)
    set_cell_margins(status_cell, top=90, start=120, bottom=90, end=120)
    set_cell_border(
        status_cell,
        top={"val": "single", "sz": 8, "color": "D4A72C"},
        bottom={"val": "single", "sz": 8, "color": "D4A72C"},
        start={"val": "single", "sz": 8, "color": "D4A72C"},
        end={"val": "single", "sz": 8, "color": "D4A72C"},
    )
    p = status_cell.paragraphs[0]
    add_text(p, data["status"], bold=True, color="633C01", size=8)
    p.paragraph_format.space_after = Pt(2)
    for notice in data["notices"]:
        notice_p = status_cell.add_paragraph()
        add_text(notice_p, notice, color="633C01", size=8.3)
        notice_p.paragraph_format.space_after = Pt(2)
        set_paragraph_rules(notice_p)

    summary = document.add_paragraph()
    summary.paragraph_format.space_before = Pt(11)
    summary.paragraph_format.space_after = Pt(12)
    add_text(summary, data["summary"], color=MUTED, size=10)
    set_paragraph_rules(summary)


def add_parties(document: Document, data: dict) -> None:
    table = document.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.8)
    values = (
        ("Company", data["parties"]["company"]),
        ("Creator", data["parties"]["creator"]),
        ("Effective date", data["parties"]["effectiveDate"]),
    )
    for index, (label, value) in enumerate(values):
        row = table.rows[index]
        keep_row_together(row)
        for cell in row.cells:
            set_cell_margins(cell, top=75, start=100, bottom=75, end=100)
            set_cell_border(cell, bottom={"val": "single", "sz": 4, "color": LINE})
        set_cell_shading(row.cells[0], SOFT)
        add_text(row.cells[0].paragraphs[0], label.upper(), bold=True, color=MUTED, size=7.3)
        add_text(row.cells[1].paragraphs[0], value, bold=True, color=INK, size=8.5)


def add_key_terms(document: Document, data: dict) -> None:
    heading = document.add_paragraph(style="Heading 1")
    add_text(heading, "Proposed key terms", bold=True, size=15)
    subtitle = document.add_paragraph()
    add_text(
        subtitle,
        "Supported economics are separated from terms that still require owner or counsel approval.",
        color=MUTED,
        size=8.5,
    )
    set_paragraph_rules(subtitle)

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.45), Inches(3.65), Inches(1.95))
    for index, width in enumerate(widths):
        table.columns[index].width = width
    header = table.rows[0]
    for index, text in enumerate(("Term", "Proposed value", "Approval state")):
        cell = header.cells[index]
        set_cell_shading(cell, DARK)
        set_cell_margins(cell, top=75, start=90, bottom=75, end=90)
        add_text(cell.paragraphs[0], text, bold=True, color=WHITE, size=7.5)

    for term in data["keyTerms"]:
        row = table.add_row()
        keep_row_together(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell in row.cells:
            set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
            set_cell_border(cell, bottom={"val": "single", "sz": 4, "color": LINE})
        add_text(row.cells[0].paragraphs[0], term["label"], bold=True, color=INK, size=7.8)
        add_text(row.cells[1].paragraphs[0], term["value"], color=INK, size=7.8)
        approval_color = AMBER if "approve" in term["status"].lower() else MUTED
        add_text(row.cells[2].paragraphs[0], term["status"], color=approval_color, size=7.2)


def add_rights_callout(document: Document) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, DARK)
    set_cell_margins(cell, top=160, start=175, bottom=160, end=175)
    heading = cell.paragraphs[0]
    add_label(heading, "Content and advertising rights", color="AFB8C1")
    title = cell.add_paragraph()
    add_text(
        title,
        "GoTall owns accepted program content and may use it in paid media.",
        bold=True,
        color=WHITE,
        size=14,
    )
    title.paragraph_format.space_after = Pt(6)
    body = cell.add_paragraph()
    add_text(
        body,
        "The proposed terms combine work-made-for-hire language with a present copyright assignment, "
        "editing and derivative rights, a publicity release, and express permission for Spark Ads, "
        "Partnership Ads, and other paid placements. Creator accounts and unrelated pre-existing "
        "content remain the creator’s property.",
        color="C9D1D9",
        size=8.6,
    )


def add_recitals(document: Document, data: dict) -> None:
    for index, recital in enumerate(data["recitals"]):
        p = document.add_paragraph()
        add_text(p, recital, color=MUTED, size=8.8, italic=True)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.right_indent = Inches(0.16)
        p.paragraph_format.space_after = Pt(4 if index < len(data["recitals"]) - 1 else 10)
        set_paragraph_rules(p)


def add_sections(document: Document, data: dict) -> None:
    for section in data["sections"]:
        heading = document.add_paragraph(style="Heading 1")
        add_text(heading, f'{section["number"]}. ', bold=True, color=GREEN, size=15)
        add_text(heading, section["title"], bold=True, color=INK, size=15)
        for paragraph_text in section["paragraphs"]:
            paragraph = document.add_paragraph()
            add_text(paragraph, paragraph_text, color=INK, size=9.15)
            set_paragraph_rules(paragraph)
        for bullet_text in section["bullets"]:
            paragraph = document.add_paragraph(style="List Bullet")
            add_text(paragraph, bullet_text, color=INK, size=9)
            set_paragraph_rules(paragraph)


def add_schedule(document: Document, data: dict) -> None:
    heading = document.add_paragraph(style="Heading 1")
    add_text(heading, data["scheduleA"]["title"], bold=True, color=INK, size=15)
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.7), Inches(3.2), Inches(2.15))
    for index, width in enumerate(widths):
        table.columns[index].width = width
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(("Term", "Proposed value", "Approval state")):
        cell = header.cells[index]
        set_cell_shading(cell, SOFT)
        set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
        set_cell_border(cell, bottom={"val": "single", "sz": 8, "color": LINE})
        add_text(cell.paragraphs[0], text, bold=True, color=MUTED, size=7.5)
    for item in data["scheduleA"]["rows"]:
        row = table.add_row()
        keep_row_together(row)
        for cell in row.cells:
            set_cell_margins(cell, top=62, start=90, bottom=62, end=90)
            set_cell_border(cell, bottom={"val": "single", "sz": 4, "color": LINE})
        add_text(row.cells[0].paragraphs[0], item["term"], bold=True, size=7.6)
        add_text(row.cells[1].paragraphs[0], item["proposedValue"], size=7.6)
        add_text(row.cells[2].paragraphs[0], item["approvalState"], color=MUTED, size=7.2)


def add_campaign_brief(document: Document, data: dict) -> None:
    heading = document.add_paragraph(style="Heading 1")
    add_text(heading, "Campaign Brief fields", bold=True, size=15)
    paragraph = document.add_paragraph()
    add_text(
        paragraph,
        "Each assignment should complete these fields rather than relying on implied posting or payment rules.",
        color=MUTED,
        size=8.5,
    )
    for field in data["campaignBriefFields"]:
        item = document.add_paragraph(style="List Bullet")
        add_text(item, field, size=9)
        set_paragraph_rules(item)


def add_references(document: Document, data: dict) -> None:
    heading = document.add_paragraph(style="Heading 1")
    add_text(heading, "Primary drafting references", bold=True, size=15)
    paragraph = document.add_paragraph()
    add_text(
        paragraph,
        "These sources informed the sample structure. They are not a substitute for legal advice.",
        color=MUTED,
        size=8.5,
    )
    for reference in data["references"]:
        item = document.add_paragraph(style="List Bullet")
        add_hyperlink(item, reference["title"], reference["url"])
        set_paragraph_rules(item)


def add_signatures(document: Document, data: dict) -> None:
    heading = document.add_paragraph(style="Heading 1")
    add_text(heading, "Signature blocks — sample only", bold=True, size=15)
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    entries = data["signatureBlocks"]
    placements = ((0, 0, entries[0]), (0, 1, entries[1]), (1, 0, entries[2]))
    table.cell(1, 0).merge(table.cell(1, 1))
    for row_index, col_index, entry in placements:
        cell = table.cell(row_index, col_index)
        set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
        set_cell_border(
            cell,
            top={"val": "single", "sz": 6, "color": LINE},
            bottom={"val": "single", "sz": 6, "color": LINE},
            start={"val": "single", "sz": 6, "color": LINE},
            end={"val": "single", "sz": 6, "color": LINE},
        )
        p = cell.paragraphs[0]
        add_text(p, entry["label"], bold=True, size=8.6)
        p.paragraph_format.space_after = Pt(18)
        for label, value in (("Name", entry["name"]), ("Title", entry["title"]), ("Date", entry["date"])):
            line = cell.add_paragraph()
            add_text(line, f"{label}: ", bold=True, color=MUTED, size=7.8)
            add_text(line, value, size=7.8)
            line.paragraph_format.space_after = Pt(7)
            set_paragraph_rules(line)


def build(output: Path) -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    document = Document()
    style_document(document)
    add_header_footer(document, data)
    add_title_block(document, data)
    add_parties(document, data)
    add_key_terms(document, data)
    add_rights_callout(document)
    add_recitals(document, data)
    add_sections(document, data)
    add_schedule(document, data)
    add_campaign_brief(document, data)
    add_references(document, data)
    add_signatures(document, data)

    core = document.core_properties
    core.title = f'{data["documentTitle"]} — {data["version"]}'
    core.subject = "Non-binding counsel-review sample"
    core.author = "GoTall"
    core.keywords = "creator agreement, content ownership, paid media, sample"
    core.comments = data["status"]

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.output.resolve())


if __name__ == "__main__":
    main()
